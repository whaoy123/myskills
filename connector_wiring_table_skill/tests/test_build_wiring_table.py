import csv
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "build_wiring_table.py"


def write_inputs(inp: Path, declare_fanout: bool = True):
    (inp / "ad_pin_net.csv").write_text(
        "BoardConnector,BoardPin,NetName\nJ2,12,SIG1\nJ2,13,SIG2\nJ2,25,GND\n",
        encoding="utf-8",
    )
    fan = "GND25" if declare_fanout else ""
    (inp / "connection_plan.csv").write_text(
        "SheetName,Order,Topology,Point1Code,Point1Node,Point2Code,Point2Node,BoardEndpoint,BoardConnector,BoardPinHint,NetName,PairGroup,CoreRole,FanoutId,WireType,LabelText,Include,Notes\n"
        "J2线缆,1,twisted_pair,DB25-F-25,,J2端子组,T01-S,1,J2,,SIG1,P01,signal,,0.5,,yes,\n"
        f"J2线缆,2,twisted_pair,DB25-F-25,,J2端子组,T01-G,1,J2,,GND,P01,return,{fan},0.5,,yes,\n"
        "J2线缆,3,twisted_pair,DB25-F-25,,J2端子组,T02-S,1,J2,,SIG2,P02,signal,,0.5,,yes,\n"
        f"J2线缆,4,twisted_pair,DB25-F-25,,J2端子组,T02-G,1,J2,,GND,P02,return,{fan},0.5,,yes,\n"
        "转接线,1,direct,X1,1,X2,A,0,,,SIG1,,normal,,0.5,SIG1,yes,普通一对一\n",
        encoding="utf-8",
    )
    (inp / "signal_catalog.csv").write_text(
        "NetName,SignalDefinition,WireType,ElectricalAttribute\nSIG1,信号1,0.5,\nSIG2,信号2,0.5,\nGND,共用地,0.5,\n",
        encoding="utf-8",
    )
    (inp / "interface_catalog.csv").write_text(
        "SheetName,BoardConnector,BoardConnectorModel,BoardGender,MatingConnectorModel,MatingGender,TerminalModel,MatesTo,Point2NodeHeader,Title,Description\n"
        "J2线缆,J2,DB25-M-25,公头,DB25-F-25,母头,,测试板,端子号,J2 对插线束,\n"
        "转接线,,,,,,,,节点号,普通转接线,\n",
        encoding="utf-8",
    )


class WiringTableTest(unittest.TestCase):
    def test_modes_fanout_labels_and_layout(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); inp, out = root / "input", root / "output"; inp.mkdir()
            write_inputs(inp)
            cp = subprocess.run([sys.executable, str(SCRIPT), str(inp), str(out)], capture_output=True, text=True)
            self.assertEqual(cp.returncode, 0, cp.stderr)
            with (out / "normalized_connections.csv").open("r", encoding="utf-8-sig", newline="") as f:
                rows = list(csv.DictReader(f))
            self.assertEqual([r["Point1Node"] for r in rows[:4]], ["12", "25", "13", "25"])
            self.assertEqual(rows[1]["LabelText"], "SIG1-GND")
            self.assertEqual(rows[3]["LabelText"], "SIG2-GND")
            wb = load_workbook(out / "wiring_table.xlsx")
            self.assertEqual(wb.sheetnames, ["J2线缆", "转接线", "接口与端子"])
            ws = wb["J2线缆"]
            self.assertEqual(str(ws.merged_cells), "A1:G1 A2:A3 B2:C2 D2:E2 F2:F3 G2:G3")
            self.assertEqual(ws["B2"].value, "连接点1")
            self.assertEqual(ws["E3"].value, "端子号")
            self.assertEqual(wb["转接线"]["E3"].value, "节点号")
            self.assertIn("标签：SIG1", ws["G4"].value)
            self.assertGreaterEqual(ws.row_dimensions[4].height, 22)
            self.assertIn("Twisted-pair groups: 2", (out / "validation_report.md").read_text(encoding="utf-8"))

    def test_duplicate_endpoint_requires_fanout_id(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); inp, out = root / "input", root / "output"; inp.mkdir()
            write_inputs(inp, declare_fanout=False)
            cp = subprocess.run([sys.executable, str(SCRIPT), str(inp), str(out)], capture_output=True, text=True)
            self.assertNotEqual(cp.returncode, 0)
            self.assertIn("FanoutId", cp.stderr)


if __name__ == "__main__":
    unittest.main()
