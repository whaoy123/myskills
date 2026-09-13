import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "build_soldering_table.py"


def write_inputs(inp: Path, compatible="R0603"):
    (inp / "bom.csv").write_text(
        "Designator,Nominal,ResolvedDescription,Footprint,DesignModel,FinalModel\n"
        "R1,10kΩ,精度:±1%;,R0603,OLD-R,R-10K-0603\n"
        "R2,10kΩ,精度:±1%;,R0603,OLD-R,R-10K-0603\n"
        "J1,DB25,DB25公头;,DB25,DB25-M,DB25-M\n"
        "C9,DNP,,C0603,CUSTOM-C,CUSTOM-C\n",
        encoding="utf-8",
    )
    (inp / "component_rules.csv").write_text(
        "MatchField,MatchValue,DisplayPackage,MountType,PinsPerPart,FixedPinsPerPart,Include,CompatibleFootprints,DNPReason,ConfirmedBy,ConfirmedDate,Notes\n"
        f"Model,R-10K-0603,0603,SMD,2,0,yes,{compatible},,,,\n"
        "Footprint,DB25,DB25公头,THT,25,2,yes,,,,,含固定脚\n"
        "Model,CUSTOM-C,0603,NONE,0,0,no,C0603,设计确认不焊,用户,2026-09-13,\n",
        encoding="utf-8",
    )
    (inp / "procurement.csv").write_text(
        "Model,StockQty,PurchasedQty,ReceivedQty,ReservedQty,TargetMultiplier,Source,Notes\n"
        "R-10K-0603,1,3,,1,1.2,旧库存+当前订单,待到货\n"
        "DB25-M,0,1,1,0,1.2,采购单,\n",
        encoding="utf-8",
    )


class SolderingTableTest(unittest.TestCase):
    def test_margin_reserve_dnp_and_fixed_layout(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); inp, out = root / "input", root / "output"; inp.mkdir()
            write_inputs(inp)
            cp = subprocess.run([sys.executable, str(SCRIPT), str(inp), str(out)], capture_output=True, text=True)
            self.assertEqual(cp.returncode, 0, cp.stderr)
            wb = load_workbook(out / "soldering_table.xlsx")
            self.assertEqual(wb.sheetnames, ["焊接清单", "备料核对"])
            ws = wb["焊接清单"]
            self.assertEqual(ws["A1"].fill.fgColor.rgb, "00FFFFFF")
            self.assertEqual(ws["A1"].font.color.rgb, "00000000")
            self.assertEqual(ws["D4"].value, 3)
            self.assertEqual(ws["E4"].value, 4)
            self.assertEqual(ws["F4"].value, 27)
            inv = wb["备料核对"]
            models = {inv.cell(r, 1).value: r for r in range(2, 4)}
            rr = models["R-10K-0603"]
            self.assertEqual(inv.cell(rr, 4).value, 3)  # ceil(2 * 1.2)
            self.assertEqual(inv.cell(rr, 9).value, 3)  # 1 + 3 - 1
            self.assertEqual(inv.cell(rr, 12).value, "采购数量覆盖，待到货清点")
            self.assertIn("不焊接核对", [inv.cell(r, 1).value for r in range(1, inv.max_row + 1)])
            report = (out / "validation_report.md").read_text(encoding="utf-8")
            self.assertIn("Excluded BOM items with DNP evidence: 1", report)
            self.assertIn("Total solder joints: 31", report)

    def test_model_package_mismatch_fails(self):
        with tempfile.TemporaryDirectory() as td:
            root = Path(td); inp, out = root / "input", root / "output"; inp.mkdir()
            write_inputs(inp, compatible="R0805")
            cp = subprocess.run([sys.executable, str(SCRIPT), str(inp), str(out)], capture_output=True, text=True)
            self.assertNotEqual(cp.returncode, 0)
            self.assertNotEqual(cp.stderr, "")


if __name__ == "__main__":
    unittest.main()
